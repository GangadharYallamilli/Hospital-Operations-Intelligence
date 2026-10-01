from fastapi import APIRouter, UploadFile, File, HTTPException, Request
from pathlib import Path
import csv
import json
import logging
import re
import shutil
import uuid
from collections import Counter

from pydantic import BaseModel

from backend.data_understanding import understand_dataset
from backend.analytics import calculate_analytics


# =========================================================
# ROUTER
# =========================================================

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_CHUNK_CHARS = 1800
CHUNK_OVERLAP = 250
MAX_SEARCH_RESULTS = 5

router = APIRouter(
    prefix="/documents",
    tags=["Documents"],
)


# =========================================================
# DIRECTORIES
# =========================================================

BASE_DIR = Path(__file__).resolve().parent

UPLOAD_DIR = BASE_DIR / "uploads"
PROCESSED_DIR = BASE_DIR / "processed"

UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

PROCESSED_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# =========================================================
# USER STORAGE
# =========================================================

def user_storage_dirs(request: Request):
    """Return private document folders for the authenticated account."""

    user_id = request.session.get("user_id")

    if user_id is None:
        raise HTTPException(
            status_code=401,
            detail="Authentication required.",
        )

    try:
        user_folder = str(int(user_id))

    except (TypeError, ValueError):
        raise HTTPException(
            status_code=401,
            detail="Authentication required.",
        )

    upload_dir = UPLOAD_DIR / user_folder
    processed_dir = PROCESSED_DIR / user_folder

    upload_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    processed_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    return upload_dir, processed_dir


# =========================================================
# REQUEST MODEL
# =========================================================

class DocumentQuestion(BaseModel):

    question: str

    document_name: str | None = None


# =========================================================
# TEXT CLEANING
# =========================================================

def clean_text(value: str) -> str:
    """Normalize extracted document text."""

    if not value:
        return ""

    value = value.replace("\x00", " ")

    value = re.sub(
        r"[ \t]+",
        " ",
        value,
    )

    value = re.sub(
        r"\n{3,}",
        "\n\n",
        value,
    )

    return value.strip()


# =========================================================
# CHUNKING
# =========================================================

def chunk_text(
    text: str,
    max_chars: int = MAX_CHUNK_CHARS,
    overlap: int = CHUNK_OVERLAP,
):
    """
    Lightweight deterministic chunking.

    No embeddings or PyTorch are required.
    """

    text = clean_text(text)

    if not text:
        return []

    paragraphs = [
        part.strip()
        for part in re.split(
            r"\n\s*\n",
            text,
        )
        if part.strip()
    ]

    chunks = []
    current = ""

    for paragraph in paragraphs:

        if len(paragraph) <= max_chars:

            candidate = (
                f"{current}\n\n{paragraph}"
                if current
                else paragraph
            )

            if len(candidate) <= max_chars:
                current = candidate
                continue

            if current:
                chunks.append(current.strip())

            tail = current[-overlap:] if current else ""
            current = (
                f"{tail}\n\n{paragraph}"
                if tail
                else paragraph
            )

            if len(current) > max_chars:
                start = 0

                while start < len(current):
                    end = start + max_chars
                    chunks.append(
                        current[start:end].strip()
                    )
                    start = max(
                        end - overlap,
                        start + 1,
                    )

                current = ""

        else:

            if current:
                chunks.append(
                    current.strip()
                )
                current = ""

            start = 0

            while start < len(paragraph):

                end = start + max_chars

                piece = paragraph[
                    start:end
                ].strip()

                if piece:
                    chunks.append(piece)

                if end >= len(paragraph):
                    break

                start = max(
                    end - overlap,
                    start + 1,
                )

    if current:
        chunks.append(
            current.strip()
        )

    return [
        chunk
        for chunk in chunks
        if len(chunk.strip()) >= 20
    ]


# =========================================================
# PDF EXTRACTION
# =========================================================

def extract_pdf(file_path: Path):
    """
    Extract PDF text page by page.

    PyMuPDF is lightweight compared with the old
    SentenceTransformer/PyTorch dependency chain.
    """

    try:
        import fitz

    except ImportError as error:
        raise RuntimeError(
            "PDF support requires PyMuPDF."
        ) from error

    pages = []

    with fitz.open(str(file_path)) as document:

        for page_index, page in enumerate(
            document,
            start=1,
        ):

            text = clean_text(
                page.get_text("text")
            )

            if not text:
                continue

            pages.append(
                {
                    "page_number": page_index,
                    "text": text,
                }
            )

    return pages


# =========================================================
# DOCX EXTRACTION
# =========================================================

def extract_docx(file_path: Path):
    """Extract DOCX paragraphs and tables."""

    try:
        from docx import Document

    except ImportError as error:
        raise RuntimeError(
            "DOCX support requires python-docx."
        ) from error

    document = Document(
        str(file_path)
    )

    sections = []

    for paragraph in document.paragraphs:

        text = clean_text(
            paragraph.text
        )

        if text:
            sections.append(
                text
            )

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):

        rows = []

        for row in table.rows:

            values = [
                clean_text(cell.text)
                for cell in row.cells
            ]

            values = [
                value
                for value in values
                if value
            ]

            if values:
                rows.append(
                    " | ".join(values)
                )

        if rows:

            sections.append(
                f"Table {table_index}\n"
                + "\n".join(rows)
            )

    full_text = "\n\n".join(
        sections
    )

    return [
        {
            "page_number": None,
            "text": full_text,
        }
    ] if full_text else []


# =========================================================
# DOCUMENT PROCESSING
# =========================================================

def process_text_document(
    file_path: Path,
    processed_dir: Path,
    document_id: str,
    file_name: str,
):
    """Extract, chunk, and store a PDF/DOCX without embeddings."""

    extension = (
        file_path.suffix.lower()
    )

    if extension == ".pdf":

        pages = extract_pdf(
            file_path
        )

    elif extension == ".docx":

        pages = extract_docx(
            file_path
        )

    else:

        raise ValueError(
            "Only PDF and DOCX files can use document RAG."
        )

    if not pages:

        raise ValueError(
            "No readable text was found in the document."
        )

    chunks = []

    chunk_id = 0

    for page in pages:

        page_text = page["text"]

        page_chunks = chunk_text(
            page_text
        )

        for chunk in page_chunks:

            chunk_id += 1

            chunks.append(
                {
                    "chunk_id": chunk_id,
                    "document_id": document_id,
                    "document_name": file_name,
                    "page_number": page.get(
                        "page_number"
                    ),
                    "text": chunk,
                }
            )

    if not chunks:

        raise ValueError(
            "The document did not contain enough readable text to create searchable content."
        )

    output_file = (
        processed_dir
        / f"{document_id}_chunks.json"
    )

    payload = {
        "document_id": document_id,
        "document_name": file_name,
        "file_type": extension.replace(
            ".",
            "",
        ),
        "chunk_count": len(chunks),
        "chunks": chunks,
    }

    with open(
        output_file,
        "w",
        encoding="utf-8",
    ) as handle:

        json.dump(
            payload,
            handle,
            ensure_ascii=False,
            indent=2,
        )

    return {
        "output_file": output_file,
        "pages": len(pages),
        "chunks": len(chunks),
    }


# =========================================================
# TOKENIZATION
# =========================================================

_STOPWORDS = {
    "the",
    "a",
    "an",
    "and",
    "or",
    "of",
    "to",
    "in",
    "on",
    "for",
    "with",
    "from",
    "is",
    "are",
    "was",
    "were",
    "what",
    "which",
    "how",
    "why",
    "when",
    "where",
    "this",
    "that",
    "these",
    "those",
    "about",
    "hospital",
    "report",
    "document",
}


def tokenize(text: str):
    words = re.findall(
        r"[A-Za-z0-9][A-Za-z0-9'-]+",
        text.lower(),
    )

    return [
        word
        for word in words
        if word not in _STOPWORDS
        and len(word) > 2
    ]


# =========================================================
# DOCUMENT SEARCH
# =========================================================

def search_chunks(
    question: str,
    chunks: list,
    top_k: int = MAX_SEARCH_RESULTS,
):
    """
    Lightweight lexical retrieval.

    This intentionally replaces the old SentenceTransformer
    retrieval for the deployment-safe version.
    """

    query_tokens = tokenize(
        question
    )

    if not query_tokens:
        return []

    query_counter = Counter(
        query_tokens
    )

    scored = []

    for chunk in chunks:

        text = str(
            chunk.get(
                "text",
                "",
            )
        )

        tokens = tokenize(text)

        if not tokens:
            continue

        token_counter = Counter(
            tokens
        )

        overlap = sum(
            min(
                query_counter[token],
                token_counter[token],
            )
            for token in query_counter
        )

        if overlap <= 0:
            continue

        unique_overlap = sum(
            1
            for token in query_counter
            if token in token_counter
        )

        phrase_bonus = 0.0

        normalized_question = " ".join(
            query_tokens
        )

        normalized_text = " ".join(
            tokens
        )

        if normalized_question and (
            normalized_question
            in normalized_text
        ):
            phrase_bonus = 2.0

        score = (
            overlap
            + (unique_overlap * 0.5)
            + phrase_bonus
        )

        scored.append(
            (
                score,
                chunk,
            )
        )

    scored.sort(
        key=lambda item: item[0],
        reverse=True,
    )

    results = []

    for score, chunk in scored[:top_k]:

        item = dict(chunk)

        item["score"] = round(
            float(score),
            4,
        )

        results.append(item)

    return results


# =========================================================
# GEMINI ANSWER
# =========================================================

def generate_gemini_answer(
    question: str,
    results: list,
):
    """
    Generate a source-grounded answer with the existing
    Gemini key manager.

    No Gemini call is made when no API key is configured.
    """

    if not results:
        return None

    try:

        from backend.ai.gemini_key_manager import (
            gemini_manager,
        )

    except Exception as error:

        logger.warning(
            "Gemini manager unavailable: %r",
            error,
        )

        return None

    context_parts = []

    for index, result in enumerate(
        results,
        start=1,
    ):

        page = result.get(
            "page_number"
        )

        page_label = (
            f"Page {page}"
            if page is not None
            else "Document section"
        )

        context_parts.append(
            f"[SOURCE {index} - {page_label}]\n"
            f"{result.get('text', '')}"
        )

    context = "\n\n".join(
        context_parts
    )

    prompt = f"""
You are MedaNexa, a hospital operations intelligence assistant.

Answer the user's question ONLY from the supplied document passages.

Do not invent facts.
Do not use outside knowledge.
If the passages do not contain enough information, say:
"The information is not available in the uploaded document."

Question:
{question}

Document passages:
{context}

Give a concise, clear hospital-operations answer.
When possible, mention the relevant source/page naturally.
"""

    try:

        answer = gemini_manager.generate(
            prompt
        )

        if answer:
            return answer.strip()

    except Exception as error:

        logger.warning(
            "Gemini document answer failed: %r",
            error,
        )

    return None


# =========================================================
# EXTRACTIVE FALLBACK
# =========================================================

def extractive_fallback(
    question: str,
    results: list,
):
    """
    Return relevant retrieved passages when Gemini is unavailable.

    This prevents document intelligence from completely failing
    because of an API quota/error.
    """

    if not results:
        return (
            "The information is not available in the uploaded document."
        )

    lines = [
        "Gemini was unavailable, so here are the most relevant "
        "passages retrieved from the uploaded document:"
    ]

    for index, result in enumerate(
        results[:3],
        start=1,
    ):

        page = result.get(
            "page_number"
        )

        if page is not None:
            source = (
                f"Source {index} "
                f"(Page {page})"
            )
        else:
            source = (
                f"Source {index}"
            )

        text = clean_text(
            result.get(
                "text",
                "",
            )
        )

        if len(text) > 700:
            text = (
                text[:700].rstrip()
                + "..."
            )

        lines.append(
            f"\n{source}:\n{text}"
        )

    return "\n".join(
        lines
    )


# =========================================================
# BUILD SOURCES
# =========================================================

def build_sources(
    results: list,
):
    sources = []

    for result in results:

        source = {
            "file_name": result.get(
                "document_name",
                "Hospital document",
            ),
            "document_name": result.get(
                "document_name",
                "Hospital document",
            ),
            "type": "Uploaded document",
            "score": result.get(
                "score",
                0,
            ),
        }

        page = result.get(
            "page_number"
        )

        if page is not None:
            source["page_number"] = page

        sources.append(
            source
        )

    return sources


# =========================================================
# UPLOAD DOCUMENT
# =========================================================

@router.post("/upload")
async def upload_document(
    request: Request,
    file: UploadFile = File(...),
):

    if not file.filename:

        raise HTTPException(
            status_code=400,
            detail="No file selected.",
        )

    extension = (
        Path(file.filename)
        .suffix
        .lower()
    )

    allowed_extensions = {
        ".pdf",
        ".docx",
        ".csv",
        ".xlsx",
    }

    if extension not in allowed_extensions:

        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported document format. "
                "Please upload a PDF, DOCX, CSV, or XLSX file."
            ),
        )

    file.file.seek(
        0,
        2,
    )

    upload_size = file.file.tell()

    file.file.seek(0)

    if upload_size > MAX_UPLOAD_BYTES:

        raise HTTPException(
            status_code=413,
            detail=(
                "Document exceeds the 25 MB upload limit."
            ),
        )

    document_id = str(
        uuid.uuid4()
    )

    upload_dir, processed_dir = (
        user_storage_dirs(request)
    )

    safe_name = Path(
        file.filename
    ).name

    stored_name = (
        f"{document_id}_{safe_name}"
    )

    file_path = (
        upload_dir
        / stored_name
    )

    try:

        # -----------------------------------------------------
        # SAVE FILE
        # -----------------------------------------------------

        with open(
            file_path,
            "wb",
        ) as buffer:

            shutil.copyfileobj(
                file.file,
                buffer,
            )

        logger.info(
            "File uploaded: %s",
            safe_name,
        )

        # -----------------------------------------------------
        # PDF / DOCX
        # -----------------------------------------------------

        if extension in {
            ".pdf",
            ".docx",
        }:

            processing = process_text_document(
                file_path,
                processed_dir,
                document_id,
                safe_name,
            )

            return {
                "status": "success",
                "document_id": document_id,
                "file_name": safe_name,
                "file_type": extension.replace(
                    ".",
                    "",
                ),
                "processing_type": "rag",
                "pages": processing["pages"],
                "chunks": processing["chunks"],
                "message": (
                    "Document uploaded and processed successfully."
                ),
            }

        # -----------------------------------------------------
        # CSV / XLSX
        # -----------------------------------------------------

        understanding = understand_dataset(
            str(file_path)
        )

        analytics = calculate_analytics(
            str(file_path)
        )

        return {
            "status": "success",
            "document_id": document_id,
            "file_name": safe_name,
            "file_type": extension.replace(
                ".",
                "",
            ),
            "processing_type": "dataset",
            "message": (
                "Dataset uploaded and analyzed successfully."
            ),
            "understanding": understanding,
            "analytics": analytics,
        }

    except HTTPException:
        raise

    except Exception as error:

        logger.exception(
            "Document processing failed"
        )

        if file_path.exists():

            try:
                file_path.unlink()

            except Exception:
                pass

        raise HTTPException(
            status_code=500,
            detail=(
                "Document processing failed: "
                f"{str(error)}"
            ),
        )


# =========================================================
# LIST DOCUMENTS
# =========================================================

@router.get("/")
def list_documents(
    request: Request,
):

    documents = []

    upload_dir, _ = (
        user_storage_dirs(request)
    )

    for file in upload_dir.iterdir():

        if not file.is_file():
            continue

        extension = (
            file.suffix.lower()
        )

        if extension not in {
            ".pdf",
            ".docx",
            ".csv",
            ".xlsx",
        }:
            continue

        processing_type = (
            "rag"
            if extension in {
                ".pdf",
                ".docx",
            }
            else "dataset"
        )

        documents.append(
            {
                "file_name": file.name,
                "size": file.stat().st_size,
                "file_type": extension.replace(
                    ".",
                    "",
                ),
                "processing_type": processing_type,
            }
        )

    documents.sort(
        key=lambda item: item["file_name"].lower()
    )

    return {
        "status": "success",
        "documents": documents,
    }


# =========================================================
# LOAD STORED CHUNKS
# =========================================================

def load_document_chunks(
    processed_dir: Path,
):
    documents = []

    for file in processed_dir.glob(
        "*_chunks.json"
    ):

        try:

            with open(
                file,
                "r",
                encoding="utf-8",
            ) as handle:

                data = json.load(
                    handle
                )

            documents.append(
                data
            )

        except Exception as error:

            logger.warning(
                "Could not read processed document %s: %r",
                file.name,
                error,
            )

    return documents


# =========================================================
# ASK DOCUMENT QUESTION
# =========================================================

@router.post("/ask")
def ask_document_question(
    payload: DocumentQuestion,
    http_request: Request,
):

    question = (
        payload.question or ""
    ).strip()

    if not question:

        raise HTTPException(
            status_code=400,
            detail="Question cannot be empty.",
        )

    _, processed_dir = (
        user_storage_dirs(
            http_request
        )
    )

    processed_documents = (
        load_document_chunks(
            processed_dir
        )
    )

    if not processed_documents:

        raise HTTPException(
            status_code=404,
            detail=(
                "No processed documents found. "
                "Please upload a PDF or DOCX document first."
            ),
        )

    # -----------------------------------------------------
    # Select document
    # -----------------------------------------------------

    selected_documents = (
        processed_documents
    )

    if payload.document_name:

        requested_name = Path(
            payload.document_name
        ).name.lower()

        selected_documents = [
            document
            for document in processed_documents
            if (
                str(
                    document.get(
                        "document_name",
                        "",
                    )
                ).lower()
                == requested_name
            )
            or requested_name.endswith(
                "_" + str(
                    document.get(
                        "document_name",
                        "",
                    )
                ).lower()
            )
        ]

        if not selected_documents:

            raise HTTPException(
                status_code=404,
                detail=(
                    f"No processed document found for "
                    f"'{payload.document_name}'."
                ),
            )

    # -----------------------------------------------------
    # Combine chunks
    # -----------------------------------------------------

    chunks = []

    for document in selected_documents:

        document_chunks = document.get(
            "chunks",
            [],
        )

        if isinstance(
            document_chunks,
            list,
        ):

            chunks.extend(
                document_chunks
            )

    # -----------------------------------------------------
    # Search
    # -----------------------------------------------------

    results = search_chunks(
        question,
        chunks,
        top_k=MAX_SEARCH_RESULTS,
    )
    summary_words = {
    "summarize",
    "summarise",
    "summary",
    "overview",
    "overall",
    }

    question_words = set(
        re.findall(
            r"[A-Za-z]+",
            question.lower(),
        )
    )

    is_summary_question = bool(
        question_words.intersection(summary_words)
    )

    if not results and is_summary_question and chunks:
        total_chunks = len(chunks)

        if total_chunks <= MAX_SEARCH_RESULTS:
            results = chunks
        else:
            indexes = [
                round(
                    i * (total_chunks - 1)
                    / (MAX_SEARCH_RESULTS - 1)
                )
                for i in range(MAX_SEARCH_RESULTS)
            ]

            results = [
                chunks[index]
                for index in indexes
            ]

    if not results:

        return {
            "status": "no_source",
            "question": question,
            "answer": (
                "The information is not available "
                "in the uploaded document."
            ),
            "sources": [],
        }

    # -----------------------------------------------------
    # Gemini
    # -----------------------------------------------------

    answer = generate_gemini_answer(
        question,
        results,
    )

    # -----------------------------------------------------
    # Local fallback
    # -----------------------------------------------------

    if not answer:

        answer = extractive_fallback(
            question,
            results,
        )

    # -----------------------------------------------------
    # Response
    # -----------------------------------------------------

    return {
        "status": "success",
        "question": question,
        "answer": answer,
        "sources": build_sources(
            results
        ),
    }
