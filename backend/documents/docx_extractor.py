import os
from docx import Document


def extract_docx(file_path):
    """
    Generic DOCX extractor.

    Automatically extracts:
    - paragraphs
    - headings
    - tables
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if not text:
            continue

        paragraphs.append({
            "text": text,
            "style": paragraph.style.name
        })

    tables = []

    for table_index, table in enumerate(
        document.tables,
        start=1
    ):

        rows = []

        for row in table.rows:

            cells = []

            for cell in row.cells:

                cells.append(
                    cell.text.strip()
                )

            rows.append(cells)

        tables.append({
            "table_number": table_index,
            "rows": rows
        })

    return {
        "document_name": os.path.basename(file_path),
        "file_type": "docx",
        "total_paragraphs": len(paragraphs),
        "total_tables": len(tables),
        "paragraphs": paragraphs,
        "tables": tables
    }


if __name__ == "__main__":

    test_folder = os.path.abspath(
        os.path.join(
            os.path.dirname(__file__),
            "uploads"
        )
    )

    for file_name in os.listdir(test_folder):

        if not file_name.lower().endswith(".docx"):
            continue

        file_path = os.path.join(
            test_folder,
            file_name
        )

        print("\nProcessing:", file_name)

        try:

            result = extract_docx(file_path)

            print(
                "Paragraphs:",
                result["total_paragraphs"]
            )

            print(
                "Tables:",
                result["total_tables"]
            )

            print("\nSample paragraphs:")

            for paragraph in result["paragraphs"][:5]:

                print(
                    f"[{paragraph['style']}] "
                    f"{paragraph['text']}"
                )

            print("\nSample table:")

            if result["tables"]:

                for row in result["tables"][0]["rows"][:5]:
                    print(row)

            else:

                print("No tables found.")

        except Exception as error:

            print("Error:", error)